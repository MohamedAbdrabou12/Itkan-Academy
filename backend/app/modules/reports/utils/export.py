from copy import deepcopy
import csv
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Dict, List, cast

import openpyxl
from fastapi import Response
from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from weasyprint import HTML

TEMPLATES_DIR = Path(__file__).parent.parent / "templates"
env = Environment(
    loader=FileSystemLoader(TEMPLATES_DIR),
    autoescape=select_autoescape(["html"]),
    trim_blocks=True,
    lstrip_blocks=True,
)


def export_pdf(
    report_data: List[Dict[str, Any]],
    file_name: str,
    date_from: str,
    date_to: str,
    title: str = "Report",
):
    if not report_data:
        template = env.get_template("empty_report.html")
        buffer = BytesIO()
        html_content = template.render(
            title=title,
            date_from=date_from,
            date_to=date_to,
        )
        HTML(string=html_content).write_pdf(buffer)
        pdf_content = buffer.getvalue()
        buffer.close()
    else:
        headers = get_headers(report_data)

        rows = []
        for record in report_data:
            row = []
            for header in headers:
                value = record.get(header, "")
                row.append(str(value))
            rows.append(row)

        template = env.get_template("report_table.html")

        generated_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        current_year = datetime.now().strftime("%Y")

        translate_headers(headers)
        html_content = template.render(
            title=title,
            date_from=date_from,
            date_to=date_to,
            generated_date=generated_date,
            headers=headers,
            rows=rows,
            total_records=len(rows),
            current_year=current_year,
        )

        buffer = BytesIO()
        HTML(string=html_content).write_pdf(buffer)
        pdf_content = buffer.getvalue()
        buffer.close()

    start_date = datetime.fromisoformat(date_from).strftime("%Y-%m-%d")
    end_date = datetime.fromisoformat(date_to).strftime("%Y-%m-%d")
    today = datetime.now().strftime("%Y-%m-%d")

    filename = f"Itkan_{file_name}_{start_date}_to_{end_date}_generated_at_{today}.pdf"

    return Response(
        content=pdf_content,
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Type": "application/pdf",
        },
    )


def export_csv(
    report_data: List[Dict[str, Any]],
    file_name: str,
    date_from: str,
    date_to: str,
) -> Response:
    if not report_data:
        csv_content = ""
    else:
        headers = get_headers(report_data)

        output = StringIO()

        # Write UTF-8 BOM for Excel compatibility (especially for Arabic)
        output.write("\ufeff")

        translated_headers = deepcopy(headers)
        translate_headers(translated_headers)

        writer = csv.writer(output)
        writer.writerow(translated_headers)

        for record in report_data:
            row = []
            for header in headers:
                value = record.get(header, "")

                if value is None:
                    row.append("")
                else:
                    row.append(str(value))
            writer.writerow(row)

        csv_content = output.getvalue()
        output.close()

    start_date = datetime.fromisoformat(date_from).strftime("%Y-%m-%d")
    end_date = datetime.fromisoformat(date_to).strftime("%Y-%m-%d")
    today = datetime.now().strftime("%Y-%m-%d")

    filename = f"Itkan_{file_name}_{start_date}_to_{end_date}_generated_at_{today}.csv"

    return Response(
        content=csv_content,
        media_type="text/csv",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Type": "text/csv; charset=utf-8-sig",
        },
    )


def export_excel(
    report_data: List[Dict[str, Any]],
    file_name: str,
    date_from: str,
    date_to: str,
) -> Response:
    wb: Workbook = openpyxl.Workbook()
    ws: Worksheet = cast(Worksheet, wb.active)

    if not report_data:
        ws.append(["No data available"])
    else:
        headers = get_headers(report_data)

        translated_headers = deepcopy(headers)
        translate_headers(translated_headers)
        ws.append(translated_headers)

        # Style header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(
            start_color="10B981", end_color="10B981", fill_type="solid"
        )

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Write data rows
        for record in report_data:
            row = []
            for header in headers:
                value = record.get(header, "")

                if value is None:
                    row.append("")
                else:
                    row.append(str(value))
            ws.append(row)

        for column in ws.iter_cols(
            min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column
        ):
            max_length = 0
            first_cell = column[0]
            if first_cell.column is not None:
                column_letter = get_column_letter(first_cell.column)
            else:
                continue

            for cell in column:
                if cell.value is not None:
                    cell_value = str(cell.value)
                    cell_length = len(cell_value)
                    if cell_length > max_length:
                        max_length = cell_length

            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        for row in range(2, ws.max_row + 1):
            if row % 2 == 0:
                fill_color = "F0FDF4"
            else:
                fill_color = "FFFFFF"

            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = PatternFill(
                    start_color=fill_color, end_color=fill_color, fill_type="solid"
                )

                cell.alignment = Alignment(horizontal="right", vertical="top")

    buffer = BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()

    start_date = datetime.fromisoformat(date_from).strftime("%Y-%m-%d")
    end_date = datetime.fromisoformat(date_to).strftime("%Y-%m-%d")
    today = datetime.now().strftime("%Y-%m-%d")

    filename = f"Itkan_{file_name}_{start_date}_to_{end_date}_generated_at_{today}.xlsx"

    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
    )


__headers_order = [
    "branch_id",
    "branch_name",
    "class_id",
    "class_name",
    "student_id",
    "student_name",
    "teacher_id",
    "teacher_name",
    "employee_id",
    "employee_name",
    "role_name",
    "classes",
    "date",
    "status",
    "check_in",
    "check_out",
    "worked_minutes",
    "evaluation_score",
    "evaluator_name",
    "total_revenue",
    "payment_count",
    "invoice_id",
    "amount",
    "due_date",
    "description",
    "record_id",
    "cycle_id",
    "cycle_name",
    "base_salary",
    "allowance",
    "bonuses",
    "deductions",
    "net_salary",
    "payment_id",
    "gateway",
    "external_txn_id",
    "paid_at",
]


def get_headers(report_data: List[Dict[str, Any]]):
    headers = []
    for record in report_data:
        for key in record.keys():
            if key not in headers:
                try:
                    headers.insert(__headers_order.index(key), key)
                except ValueError:
                    headers.append(key)

    return headers


__headers_translation = {
    "branch_id": "رقم الفرع",
    "branch_name": "اسم الفرع",
    "class_id": "رقم الفصل",
    "class_name": "اسم الفصل",
    "student_id": "رقم الطالب",
    "student_name": "اسم الطالب",
    "date": "اليوم",
    "employee_id": "رقم الموظف",
    "employee_name": "اسم الموظف",
    "teacher_id": "رقم المعلم",
    "teacher_name": "اسم المعلم",
    "role_name": "الوظيفة",
    "classes": "الفصول الدراسية",
    "status": "حالة الحضور",
    "check_in": "وقت الحضور",
    "check_out": "وقت الانصراف",
    "worked_minutes": "دقائق العمل",
    "evaluation_score": "التقييم",
    "evaluator_name": "المقيّم",
    "total_revenue": "إجمالي الإيرادات",
    "payment_count": "عدد العمليات",
    "invoice_id": "رقم الفاتورة",
    "amount": "المبلغ",
    "due_date": "تاريخ الاستحقاق",
    "description": "الوصف",
    "record_id": "رقم السجل",
    "cycle_id": "رقم الدورة",
    "cycle_name": "دورة الرواتب",
    "base_salary": "الراتب الأساسي",
    "allowance": "البدلات",
    "bonuses": "المكافآت",
    "deductions": "الاستقطاعات",
    "net_salary": "صافي الراتب",
    "payment_id": "رقم عملية الدفع",
    "gateway": "بوابة الدفع",
    "external_txn_id": "رقم المعاملة الخارجي",
    "paid_at": "تاريخ الدفع",
}


def translate_headers(headers: List[str]):
    for i, header in enumerate(headers):
        if header in __headers_translation:
            headers[i] = __headers_translation[header]


